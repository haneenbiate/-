"""Third-annotator adjudication for Study v3: build the workbook, then merge its decisions.

    python -m tools.build_adjudication_workbook build --package runs/study-v3/annotation-v1 \
        --a A.xlsx --b B.xlsx --output adjudication/        # workbook + private key
    python -m tools.build_adjudication_workbook finalize --package runs/study-v3/annotation-v1 \
        --a A.xlsx --b B.xlsx --adjudication returned.xlsx --key adjudication/adjudication_key_private.csv \
        --output runs/study-v3/annotation-v1-final

Only the two fields used for meaning-preservation claims are adjudicated. The two earlier
answers are shown anonymously, in a per-item random order, without the annotators' notes.
"""
import argparse
import hashlib
import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from nasn.annotation_v3 import FIELDS, read_sheet

CORE = ['meaning_preserved', 'sentiment_preserved']
ARABIC = {'yes': 'نعم', 'no': 'لا', 'unsure': 'غير متأكد'}
COLUMNS = {'item': 'رقم البند', 'original': 'النص الأصلي', 'modified': 'النص المعدّل',
           'first_meaning': 'رأي 1: المعنى', 'first_sentiment': 'رأي 1: المشاعر',
           'second_meaning': 'رأي 2: المعنى', 'second_sentiment': 'رأي 2: المشاعر',
           'meaning_preserved': 'قرارك: هل بقي المعنى؟', 'sentiment_preserved': 'قرارك: هل بقيت المشاعر؟',
           'notes': 'ملاحظات'}
INSTRUCTIONS = [
    ('تعليمات الحسم — اقرأها قبل البدء', True),
    ('', False),
    ('اختلف مراجعان سابقان في هذه البنود حول سؤالين فقط. مهمتك أن تقرر الإجابة الصحيحة لكل بند.', False),
    ('يظهر لك رأياهما («رأي 1» و«رأي 2») للاطلاع فقط، وترتيبهما عشوائي. لك أن توافق أحدهما أو تختار إجابة أخرى.', False),
    ('اقرأ التغريدة الأصلية والمعدّلة كاملتين، ثم اختر من القوائم المنسدلة في العمودين الملونين.', False),
    ('', False),
    ('السؤالان', True),
    ('هل بقي المعنى؟ (yes / no / unsure): هل يفهم القارئ العربي النص المعدّل بنفس معنى الأصل؟', False),
    ('هل بقيت المشاعر؟ (yes / no / unsure): هل بقي اتجاه المشاعر وشدتها كما هي تقريبًا؟', False),
    ('', False),
    ('قواعد عامة', True),
    ('• احكم على التغريدة كاملة، لا على الكلمة المتغيرة وحدها.', False),
    ('• التوكيد بتكرار الحرف («حلووو») قد يزيد الشدة؛ إذا تغيّرت الشدة بوضوح فالمشاعر لم تبقَ (no).', False),
    ('• إذا صارت الكلمة كلمة أخرى موجودة في اللغة فالمعنى لم يبقَ (no).', False),
    ('• استخدم unsure عند الشك الحقيقي فقط.', False),
    ('', False),
    ('الوقت المتوقع: ساعة إلى ساعتين. عند الانتهاء أرسل الملف إلى المنسّق.', False),
]
INPUT = PatternFill('solid', fgColor='EEF4FC')
HEADER = PatternFill('solid', fgColor='2A78D6')


def disagreements(package, path_a, path_b):
    a, b = read_sheet(path_a, package), read_sheet(path_b, package)
    differ = (a[CORE] != b[CORE]).any(axis=1)
    return a, b, differ


def build(package, path_a, path_b, output):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    a, b, differ = disagreements(package, path_a, path_b)
    rows, key = [], []
    for i in differ[differ].index:
        # Deterministic per-item coin flip decides which annotator is shown first.
        a_first = int(hashlib.sha256(f'42|adjudication|{a.item[i]}'.encode()).hexdigest(), 16) % 2 == 0
        first, second = (a, b) if a_first else (b, a)
        rows.append({'item': a.item[i], 'original': a.original[i], 'modified': a.modified[i],
                     'first_meaning': ARABIC[first.meaning_preserved[i]], 'first_sentiment': ARABIC[first.sentiment_preserved[i]],
                     'second_meaning': ARABIC[second.meaning_preserved[i]], 'second_sentiment': ARABIC[second.sentiment_preserved[i]]})
        key.append({'item': a.item[i], 'first': 'A' if a_first else 'B'})
    pd.DataFrame(key).to_csv(output / 'adjudication_key_private.csv', index=False, lineterminator='\n')

    book = Workbook()
    guide = book.active
    guide.title = 'التعليمات'
    guide.sheet_view.rightToLeft = True
    guide.column_dimensions['A'].width = 120
    for r, (text, bold) in enumerate(INSTRUCTIONS, 1):
        cell = guide.cell(r, 1, text)
        cell.font = Font(bold=bold, size=13 if bold else 11)
        cell.alignment = Alignment(wrap_text=True, vertical='top', horizontal='right')
    guide.cell(len(INSTRUCTIONS) + 2, 1, f'عدد البنود: {len(rows)}').font = Font(bold=True)

    sheet = book.create_sheet('الحسم')
    sheet.sheet_view.rightToLeft = True
    names = list(COLUMNS)
    widths = {'item': 10, 'original': 50, 'modified': 50, 'notes': 30}
    for j, name in enumerate(names, 1):
        cell = sheet.cell(1, j, COLUMNS[name])
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = HEADER
        cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
        sheet.column_dimensions[cell.column_letter].width = widths.get(name, 14)
    for r, row in enumerate(rows, 2):
        for j, name in enumerate(names, 1):
            cell = sheet.cell(r, j, row.get(name))
            cell.alignment = Alignment(wrap_text=True, vertical='top', horizontal='right')
            if name in CORE or name == 'notes':
                cell.fill = INPUT
            else:
                cell.number_format = '@'
    for j, name in enumerate(names, 1):
        if name in CORE:
            letter = sheet.cell(1, j).column_letter
            rule = DataValidation(type='list', formula1='"' + ','.join(FIELDS[name]) + '"', allow_blank=True,
                                  showErrorMessage=True, errorTitle='قيمة غير مسموحة',
                                  error='اختر من القائمة: ' + ' / '.join(FIELDS[name]))
            rule.add(f'{letter}2:{letter}{len(rows) + 1}')
            sheet.add_data_validation(rule)
    sheet.freeze_panes = 'B2'
    book.active = 1
    book.save(output / 'توسيم_المحكّم.xlsx')
    return len(rows)


def read_adjudication(path):
    """The returned workbook, either as built here or re-saved as one sheet with Arabic answers."""
    sheets = pd.read_excel(path, sheet_name=None, dtype=str)
    name = 'الحسم' if 'الحسم' in sheets else next(n for n, s in sheets.items() if COLUMNS['item'] in s.columns)
    sheet = sheets[name].fillna('')
    sheet = sheet.rename(columns={v: k for k, v in COLUMNS.items()})
    answers = {arabic: english for english, arabic in ARABIC.items()}
    for name in CORE:
        sheet[name] = sheet[name].str.strip().replace(answers)
        bad = sheet[~sheet[name].isin(FIELDS[name])]
        if not bad.empty:
            raise ValueError(f'{path}: {name} missing/invalid for {bad.item.head(5).tolist()}')
    return sheet


def finalize(package, path_a, path_b, adjudication, key_path, output):
    """Final labels: agreed answers kept, disagreements replaced by the adjudicator's decision."""
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    a, b, differ = disagreements(package, path_a, path_b)
    decided = read_adjudication(adjudication).set_index('item')
    key = pd.read_csv(key_path).set_index('item')
    expected = set(a.item[differ])
    if set(decided.index) != expected or set(key.index) != expected:
        raise ValueError('Adjudicated items do not match the current disagreements')
    variants = pd.read_csv(Path(package) / 'annotation_key_private.csv', encoding='utf-8-sig')
    final = variants[['item', 'variant_id', 'condition']].copy()
    for name in CORE:
        final[name] = [decided.loc[item, name] if item in expected else a[name][i] for i, item in enumerate(a.item)]
    final['source'] = ['adjudicated' if item in expected else 'agreed' for item in a.item]
    final['preserved'] = final.meaning_preserved.eq('yes') & final.sentiment_preserved.eq('yes')
    final.to_csv(output / 'final_decisions.csv', index=False, lineterminator='\n')
    sub = final[final.source == 'adjudicated'].set_index('item')
    agreement = {name: {'adjudicator_matches_A': float((sub[name] == a.set_index('item').loc[sub.index, name]).mean()),
                        'adjudicator_matches_B': float((sub[name] == b.set_index('item').loc[sub.index, name]).mean())}
                 for name in CORE}
    summary = {'items': len(final), 'adjudicated': int(differ.sum()), 'agreement_with_adjudicator': agreement,
               'preserved_rate_by_condition': final.groupby('condition').preserved.mean().round(4).to_dict(),
               'preserved_count_by_condition': final.groupby('condition').preserved.sum().astype(int).to_dict()}
    (output / 'final_summary.json').write_text(json.dumps(summary, indent=2, ensure_ascii=False) + '\n',
                                               encoding='utf-8', newline='\n')
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest='command', required=True)
    for name in ['build', 'finalize']:
        p = sub.add_parser(name)
        p.add_argument('--package', required=True)
        p.add_argument('--a', required=True)
        p.add_argument('--b', required=True)
        p.add_argument('--output', required=True)
        if name == 'finalize':
            p.add_argument('--adjudication', required=True)
            p.add_argument('--key', required=True)
    args = parser.parse_args()
    if args.command == 'build':
        print(build(args.package, args.a, args.b, args.output), 'items')
    else:
        print(json.dumps(finalize(args.package, args.a, args.b, args.adjudication, args.key, args.output),
                         indent=2, ensure_ascii=False))


if __name__ == '__main__':
    main()
