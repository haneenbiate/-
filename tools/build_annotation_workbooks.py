"""One self-contained Excel workbook per Study v3 annotator: instructions, items with dropdowns, declaration."""
import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from nasn.annotation_v3 import FIELDS, PILOT_ITEMS

HEADERS = {'item': 'رقم البند', 'original': 'النص الأصلي', 'modified': 'النص المعدّل',
           'meaning_preserved': 'هل بقي المعنى؟', 'sentiment_preserved': 'هل بقيت المشاعر؟',
           'modified_sentiment': 'مشاعر النص المعدّل', 'negation_present': 'هل يوجد نفي؟',
           'negation_affected': 'هل تأثر النفي؟', 'notes': 'ملاحظات'}
INSTRUCTIONS = [
    ('تعليمات التوسيم — اقرأها كاملة قبل البدء', True),
    ('', False),
    ('في ورقة «التوسيم» تجد 330 بندًا. كل بند فيه تغريدة أصلية ونسخة غُيّر فيها حرف أو أكثر.', False),
    ('أجب عن الأعمدة الملونة باختيار القيمة من القائمة المنسدلة فقط. لا تعدّل أعمدة النص.', False),
    ('', False),
    ('الأعمدة المطلوبة', True),
    ('هل بقي المعنى؟ (yes / no / unsure): هل يفهم القارئ العربي النص المعدّل بنفس معنى الأصل؟', False),
    ('هل بقيت المشاعر؟ (yes / no / unsure): هل بقي اتجاه المشاعر وشدتها كما هي تقريبًا؟', False),
    ('مشاعر النص المعدّل (NEG سلبي / NEU محايد / POS إيجابي): كما تفهمها أنت.', False),
    ('هل يوجد نفي؟ (yes / no): هل في النص نفي لغوي؟', False),
    ('هل تأثر النفي؟ (yes / no / na): هل وقع التغيير على أداة النفي أو غيّر فهمك للنفي؟ اختر na إذا لم يوجد نفي.', False),
    ('ملاحظات: اختيارية. اكتب سبب أي unsure أو no بجملة قصيرة.', False),
    ('', False),
    ('قواعد عامة', True),
    ('• احكم على التغريدة كاملة، لا على الكلمة المتغيرة وحدها.', False),
    ('• التوكيد بتكرار الحرف («حلووو») قد يزيد الشدة؛ إذا تغيّرت الشدة بوضوح فالمشاعر لم تبقَ (no).', False),
    ('• إذا صارت الكلمة كلمة أخرى موجودة في اللغة فالمعنى لم يبقَ (no).', False),
    ('• استخدم unsure عند الشك الحقيقي فقط، ولا تخمّن.', False),
    ('', False),
    ('ما يُعدّ نفيًا', True),
    ('يُعدّ نفيًا: لا، لم، لن، ليس وأخواتها، بدون، غير/سوى بعد نفي، مش، مو، موب، مب، ماكو، مافي، مفيش، و«ما» النافية.', False),
    ('لا يُعدّ نفيًا: «ما» الموصولة أو الاستفهامية أو التعجبية («ما شاء الله»)، «ولا» بمعنى «أو»، «لا زال» بمعنى الاستمرار.', False),
    ('', False),
    ('طريقة العمل', True),
    (f'1. أنجز أول {PILOT_ITEMS} بندًا (المظللة بالأصفر) وحدك، ثم تجتمعان مع المنسّق لمناقشة فهم التعليمات فقط.', False),
    ('2. أكمل باقي البنود وحدك دون أي تواصل مع المراجع الآخر حول البنود.', False),
    ('3. لا تستخدم أي أداة ذكاء اصطناعي (ChatGPT أو غيره) أو ترجمة آلية.', False),
    ('4. املأ ورقة «الإقرار»، ثم أرسل الملف من بريدك مباشرة إلى المنسّق. لا تعدّله بعد الإرسال.', False),
    ('', False),
    ('الوقت المتوقع: 4 إلى 6 ساعات، ويُستحسن توزيعه على أكثر من جلسة.', False),
]
DECLARATION = ['الاسم الكامل', 'الجهة / الجامعة', 'البريد الإلكتروني', 'اللهجة الأم', 'المؤهل العلمي',
               'تاريخ البدء', 'تاريخ الإرسال',
               'أقرّ بأنني أنجزت هذا التوسيم بنفسي، دون استخدام أدوات ذكاء اصطناعي، ودون الاطلاع على إجابات المراجع الآخر أو نتائج أي نموذج (اكتب: أقرّ)',
               'التوقيع (اكتب اسمك)']
INPUT = PatternFill('solid', fgColor='EEF4FC')
PILOT = PatternFill('solid', fgColor='FFF4CC')
HEADER = PatternFill('solid', fgColor='2A78D6')


def build(sheet_csv, output, annotator):
    items = pd.read_csv(sheet_csv, dtype=str, keep_default_na=False, encoding='utf-8-sig')
    book = Workbook()
    guide = book.active
    guide.title = 'التعليمات'
    guide.sheet_view.rightToLeft = True
    guide.column_dimensions['A'].width = 120
    for i, (text, bold) in enumerate(INSTRUCTIONS, 1):
        cell = guide.cell(i, 1, text)
        cell.font = Font(bold=bold, size=13 if bold else 11)
        cell.alignment = Alignment(wrap_text=True, vertical='top', horizontal='right')
    guide.cell(len(INSTRUCTIONS) + 2, 1, f'رمز المراجع: {annotator}').font = Font(bold=True)

    sheet = book.create_sheet('التوسيم')
    sheet.sheet_view.rightToLeft = True
    columns = list(HEADERS)
    for j, key in enumerate(columns, 1):
        cell = sheet.cell(1, j, HEADERS[key])
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = HEADER
        cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
    widths = {'item': 10, 'original': 55, 'modified': 55, 'notes': 30}
    for j, key in enumerate(columns, 1):
        sheet.column_dimensions[sheet.cell(1, j).column_letter].width = widths.get(key, 16)
    for i, row in enumerate(items.itertuples(index=False), 2):
        for j, key in enumerate(columns, 1):
            cell = sheet.cell(i, j, getattr(row, key) if key in ('item', 'original', 'modified') else None)
            cell.alignment = Alignment(wrap_text=True, vertical='top', horizontal='right')
            if key in ('item', 'original', 'modified'):
                cell.number_format = '@'
                if i - 2 < PILOT_ITEMS:
                    cell.fill = PILOT
            else:
                cell.fill = INPUT
    last = len(items) + 1
    for j, key in enumerate(columns, 1):
        if key in FIELDS:
            letter = sheet.cell(1, j).column_letter
            rule = DataValidation(type='list', formula1='"' + ','.join(FIELDS[key]) + '"', allow_blank=True,
                                  showErrorMessage=True, errorTitle='قيمة غير مسموحة',
                                  error='اختر من القائمة: ' + ' / '.join(FIELDS[key]))
            rule.add(f'{letter}2:{letter}{last}')
            sheet.add_data_validation(rule)
    sheet.freeze_panes = 'B2'

    form = book.create_sheet('الإقرار')
    form.sheet_view.rightToLeft = True
    form.column_dimensions['A'].width = 70
    form.column_dimensions['B'].width = 45
    for i, label in enumerate(DECLARATION, 1):
        form.cell(i, 1, label).alignment = Alignment(wrap_text=True, horizontal='right', vertical='top')
        form.cell(i, 2).fill = INPUT
    book.active = 1
    book.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--package', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    Path(args.output).mkdir(parents=True, exist_ok=True)
    for annotator in ['A', 'B']:
        build(Path(args.package) / f'annotation_sheet_{annotator}.csv',
              Path(args.output) / f'توسيم_المراجع_{annotator}.xlsx', annotator)


if __name__ == '__main__':
    main()
