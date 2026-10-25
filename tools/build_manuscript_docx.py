"""Build the Springer (JKSU-CIS) Word submission from the released analysis bundles.

Writes an anonymized manuscript, a separate title page, a supplementary-information file and
the figure files. Every number in a table is read from the bundles; the prose mirrors
paper/main.tex with author-year citations.

    python tools/build_manuscript_docx.py --output ../submission_JKSU-CIS
"""
import argparse
import io
import json
import re
import shutil
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from tools.report_study_v3 import CONDITION_LABELS, SHORT, clean_performance, load_bundle, ordered

ROOT = Path(__file__).resolve().parents[1]
VALIDATION = ROOT / 'reports/study-v3/results/20261024T072801Z-public.zip'
TEST = ROOT / 'reports/study-v3/results/20261025T072918Z-test-public.zip'
HUMAN = ROOT / 'paper/generated/human_validation.json'
HYPOTHESES = {'Validation': ROOT / 'paper/generated/hypotheses.csv', 'Test': ROOT / 'paper/generated-test/hypotheses.csv'}
FIGURES = {1: ROOT / 'paper/figures/tokenizer_edits', 2: ROOT / 'paper/generated-test/flip_rates'}
DATASETS = {'arsarcasm': 'ArSarcasm-v2', 'arsas': 'ArSAS'}
ARABIC = re.compile(r'[؀-ۿݐ-ݿ]+(?:[\sـ]*[؀-ۿݐ-ݿ]+)*')
TITLE = 'When one character erases a word: orthographic noise, tokenizers and negation in Arabic sentiment classification'
PLAIN = {'\\%': '%', '$\\times$': '×', '$\\geq$': '≥'}


def label(condition):
    text = CONDITION_LABELS[condition]
    for a, b in PLAIN.items():
        text = text.replace(a, b)
    return text


def pct(x, digits=1):
    return '–' if pd.isna(x) else f'{100 * x:.{digits}f}'


def ci(row, name):
    return f"{pct(row[name])} [{pct(row[name + '_low'])}, {pct(row[name + '_high'])}]"


# ---------------------------------------------------------------- document helpers

def base_document():
    doc = Document()
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    normal.element.rPr.rFonts.set(qn('w:cs'), 'Times New Roman')
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in [('Heading 1', 12), ('Heading 2', 10.5), ('Heading 3', 10)]:
        style = styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = name == 'Heading 3'
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.rPr.rFonts
        for key in ['w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme']:
            fonts.attrib.pop(qn(key), None)
        style.paragraph_format.space_before = Pt(10 if name == 'Heading 1' else 6)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
    for name in ['List Bullet']:
        styles[name].font.name = 'Times New Roman'
        styles[name].font.size = Pt(10)
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    for side in ['left_margin', 'right_margin']:
        setattr(section, side, Cm(2.3))
    section.top_margin = section.bottom_margin = Cm(2.2)
    add_page_number(section)
    return doc


def add_page_number(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    for tag, text in [('begin', None), (None, 'PAGE'), ('end', None)]:
        if tag:
            element = OxmlElement('w:fldChar')
            element.set(qn('w:fldCharType'), tag)
        else:
            element = OxmlElement('w:instrText')
            element.set(qn('xml:space'), 'preserve')
            element.text = text
        run._r.append(element)


def add_line_numbers(section):
    numbers = OxmlElement('w:lnNumType')
    numbers.set(qn('w:countBy'), '1')
    numbers.set(qn('w:restart'), 'continuous')
    numbers.set(qn('w:distance'), '284')
    section._sectPr.append(numbers)


def add_runs(paragraph, text, size=None, bold=False, italic=False):
    """Inline markup: **bold**, *italic*, [[author placeholder]]; Arabic spans get right-to-left runs."""
    for part in re.split(r'(\*\*.+?\*\*|\*.+?\*|\[\[.+?\]\])', text):
        if not part:
            continue
        b, i = bold, italic
        if part.startswith('[['):
            run_text(paragraph, '[' + part[2:-2] + ']', size, b, i).font.highlight_color = 7  # yellow
            continue
        if part.startswith('**'):
            part, b = part[2:-2], True
        elif part.startswith('*'):
            part, i = part[1:-1], True
        cursor = 0
        for match in ARABIC.finditer(part):
            if match.start() > cursor:
                run_text(paragraph, part[cursor:match.start()], size, b, i)
            run_text(paragraph, match.group(), size, b, i, arabic=True)
            # A left-to-right mark keeps neutral punctuation between Arabic spans in the English order.
            run_text(paragraph, '‎', size, b, i)
            cursor = match.end()
        if cursor < len(part):
            run_text(paragraph, part[cursor:], size, b, i)
    return paragraph


def run_text(paragraph, text, size, bold, italic, arabic=False):
    run = paragraph.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    if arabic:
        properties = run._r.get_or_add_rPr()
        fonts = properties.find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            properties.insert(0, fonts)
        for key in ['w:ascii', 'w:hAnsi', 'w:cs']:
            fonts.set(qn(key), 'Times New Roman')
        properties.append(OxmlElement('w:rtl'))
        size_cs = OxmlElement('w:szCs')
        size_cs.set(qn('w:val'), str(int(round((size or 10) * 2 + 2))))
        properties.append(size_cs)
    return run


def para(doc, text, align=None, size=None, bold=False, italic=False, after=None):
    paragraph = doc.add_paragraph()
    if align == 'center':
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align != 'left':
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if after is not None:
        paragraph.paragraph_format.space_after = Pt(after)
    return add_runs(paragraph, text, size, bold, italic)


def bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style='List Bullet')
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.space_after = Pt(2)
        add_runs(paragraph, item)


def heading(doc, text, level):
    doc.add_heading(text, level=level)


def set_cell_border(cell, **edges):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        properties.append(borders)
    for edge, value in edges.items():
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), value)
        element.set(qn('w:sz'), '6')
        element.set(qn('w:color'), '000000')
        borders.append(element)


def table_style(doc, size):
    """Paragraph style for table cells, so empty cells take the same height as filled ones."""
    name = f'Table text {size}'
    if name not in [s.name for s in doc.styles]:
        style = doc.styles.add_style(name, 1)
        style.base_style = doc.styles['Normal']
        style.font.size = Pt(size)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.line_spacing = 1.0
    return name


def table(doc, number, caption, header, rows, widths=None, size=8.5, groups=None, note=None, left=1):
    """Three-line (booktabs-like) table with the caption above, Springer style."""
    cap = doc.add_paragraph()
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.space_before = Pt(6)
    add_runs(cap, f'**Table {number}** {caption}', size=9)
    head_rows = ([groups] if groups else []) + [header]
    grid = doc.add_table(rows=len(head_rows) + len(rows), cols=len(header))
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = widths is None
    if widths:
        layout = OxmlElement('w:tblLayout')
        layout.set(qn('w:type'), 'fixed')
        grid._tbl.tblPr.append(layout)
    style = table_style(doc, size)
    together = len(rows) <= 20
    for r, values in enumerate(head_rows + rows):
        row_properties = grid.rows[r]._tr.get_or_add_trPr()
        row_properties.append(OxmlElement('w:cantSplit'))
        for c, value in enumerate(values):
            cell = grid.cell(r, c)
            cell.text = ''
            paragraph = cell.paragraphs[0]
            paragraph.style = doc.styles[style]
            paragraph.paragraph_format.keep_with_next = together and r < len(head_rows) + len(rows) - 1
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if c < left else WD_ALIGN_PARAGRAPH.CENTER
            add_runs(paragraph, str(value), size=size, bold=False)
            if widths:
                cell.width = Cm(widths[c])
            if r == 0:
                set_cell_border(cell, top='single')
            if r == len(head_rows) - 1:
                set_cell_border(cell, bottom='single')
            if r == len(head_rows) + len(rows) - 1:
                set_cell_border(cell, bottom='single')
    if groups:
        merge_groups(grid, groups)
    if note:
        para(doc, note, size=8, after=6)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return grid


def merge_groups(grid, groups):
    """Merge runs of identical non-empty labels in the group header row."""
    c = 0
    while c < len(groups):
        end = c
        while end + 1 < len(groups) and groups[end + 1] == groups[c] and groups[c]:
            end += 1
        if end > c:
            text = groups[c]
            merged = grid.cell(0, c).merge(grid.cell(0, end))
            merged.text = ''
            paragraph = merged.paragraphs[0]
            paragraph.style = grid.cell(1, 0).paragraphs[0].style
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(paragraph, text, size=8.5)
        c = end + 1


def figure(doc, number, path, caption, width=16.0):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Cm(width))
    cap = doc.add_paragraph()
    cap.paragraph_format.space_after = Pt(8)
    add_runs(cap, f'**Fig. {number}** {caption}', size=9)


# ---------------------------------------------------------------- data

class Data:
    def __init__(self):
        self.validation = load_bundle(VALIDATION)
        self.test = load_bundle(TEST)
        self.human = json.loads(HUMAN.read_text(encoding='utf-8'))
        self.hypotheses = {k: pd.read_csv(v) for k, v in HYPOTHESES.items()}

    def cells(self, bundle, dataset):
        return bundle[f'{dataset}-all-robustness_cells.csv']

    def natural(self, bundle, dataset):
        return bundle[f'{dataset}-all-natural_noise_cells.csv']

    def models(self):
        return ordered(self.cells(self.test, 'arsarcasm').model)


def clean_table(data):
    validation = clean_performance(data.validation).groupby(['dataset', 'model']).macro_f1.agg(['mean', 'std'])
    rows = []
    for model in data.models():
        row = [SHORT.get(model, model)]
        for dataset in DATASETS:
            v = validation.loc[(dataset, model)]
            row.append(f"{100 * v['mean']:.1f} ± {100 * v['std']:.1f}")
        for dataset in DATASETS:
            cells = data.cells(data.test, dataset)
            c = cells[(cells.model == model) & (cells.defense == 'RAW') & (cells.condition == 'tatweel/single/non_cue')].iloc[0]
            row.append(ci(c, 'clean_macro_f1'))
        rows.append(row)
    groups = ['', 'Validation', 'Validation', 'Test', 'Test']
    return groups, ['Model'] + list(DATASETS.values()) * 2, rows


def drop_table(bundle, data):
    conditions = ['tatweel/dense/non_cue', 'elongationx2/dense/non_cue', 'typo/dense/non_cue']
    names = ['Tatweel', 'Elong. ×2', 'Typo']
    rows = []
    for model in data.models():
        row = [SHORT.get(model, model)]
        for dataset in DATASETS:
            cells = data.cells(bundle, dataset).set_index(['model', 'defense', 'condition'])
            for condition in conditions:
                row.append(ci(cells.loc[(model, 'RAW', condition)], 'macro_f1_drop'))
        rows.append(row)
    groups = [''] + [DATASETS[d] for d in DATASETS for _ in conditions]
    return groups, ['Model'] + names * 2, rows


def flip_table(bundle, data, dataset):
    cells = data.cells(bundle, dataset).set_index(['model', 'defense', 'condition'])
    models = data.models()
    rows = [[label(c)] + [pct(cells.loc[(m, 'RAW', c), 'flip_rate']) for m in models] for c in CONDITION_LABELS]
    return ['Condition'] + [SHORT.get(m, m) for m in models], rows


def defense_table(bundle, data, defenses):
    conditions = {'tatweel/dense/non_cue': 'Tatweel', 'elongationx2/dense/non_cue': 'Elongation ×2', 'fold/dense/non_cue': 'Folding'}
    names = {'RAW': 'RAW', 'STD': 'STD', 'NO_CONTEXT': 'NO_CTX', 'NASN2': 'NASN2'}
    rows = []
    for condition, name in conditions.items():
        for j, model in enumerate(data.models()):
            row = [name if j == 0 else '', SHORT.get(model, model)]
            for dataset in DATASETS:
                cells = data.cells(bundle, dataset).set_index(['model', 'defense', 'condition'])
                row += [pct(cells.loc[(model, d, condition), 'flip_rate']) for d in defenses]
            rows.append(row)
    groups = ['', ''] + [DATASETS[d] for d in DATASETS for _ in defenses]
    return groups, ['Edit', 'Model'] + [names[d] for d in defenses] * 2, rows


def natural_table(data):
    rows = []
    header = ['Model']
    groups = ['']
    for dataset in DATASETS:
        f = data.natural(data.test, dataset)
        for noise, name in [('tatweel', 'Tatweel'), ('letter_run_ge3', 'Letter run ≥3')]:
            header.append(f"{name} (n = {int(f[f.noise == noise].n.max())})")
            groups.append(DATASETS[dataset])
    for model in data.models():
        row = [SHORT.get(model, model)]
        for dataset in DATASETS:
            f = data.natural(data.test, dataset)
            for noise in ['tatweel', 'letter_run_ge3']:
                r = f[(f.model == model) & (f.noise == noise)].iloc[0]
                row.append(ci(r, 'flip_rate'))
        rows.append(row)
    return groups, header, rows


def agreement_table(data):
    names = {'meaning_preserved': 'Meaning preserved', 'sentiment_preserved': 'Sentiment preserved',
             'modified_sentiment': 'Sentiment of edited text', 'negation_present': 'Negation present',
             'negation_affected': 'Negation affected by edit'}
    rows = [[names[r['field']], f"{100 * r['agreement']:.1f}%", f"{r['cohen_kappa']:.2f}", f"{r['gwet_ac1']:.2f}"]
            for r in data.human['agreement']]
    return ['Field', 'Agreement', 'Cohen’s κ', 'Gwet’s AC1'], rows


def preservation_table(data):
    rows = []
    for condition in CONDITION_LABELS:
        s = data.human['by_condition'][condition]
        rows.append([label(condition), f"{s['meaning_a']} / {s['meaning_b']}", s['both_meaning'], s['both_preserved'],
                     s['final_preserved']])
    return ['Condition', 'Meaning (A / B)', 'Both: meaning', 'Both: meaning + sentiment', 'Final'], rows


def hypotheses_table(data):
    rows = []
    for hypothesis in sorted(set(data.hypotheses['Validation'].hypothesis) | set(data.hypotheses['Test'].hypothesis)):
        row = [hypothesis]
        for partition in ['Validation', 'Test']:
            frame = data.hypotheses[partition]
            for dataset in DATASETS:
                part = frame[(frame.hypothesis == hypothesis) & (frame.dataset == dataset)]
                row.append('–' if part.empty else f"{int(part.verdict.eq('supported').sum())}/{len(part)}")
        rows.append(row)
    return ['', 'Validation', 'Validation', 'Test', 'Test'], ['Hypothesis'] + list(DATASETS.values()) * 2, rows


# ---------------------------------------------------------------- references (Springer author-year)

REFERENCES = [
    'Abdul-Mageed M, Elmadany A, Nagoudi EMB (2021) ARBERT & MARBERT: deep bidirectional transformers for Arabic. In: Proceedings of the 59th Annual Meeting of the Association for Computational Linguistics and the 11th International Joint Conference on Natural Language Processing (Volume 1: Long Papers). Association for Computational Linguistics, pp 7088–7105. https://doi.org/10.18653/v1/2021.acl-long.551',
    'Abu Farha I, Zaghouani W, Magdy W (2021) Overview of the WANLP 2021 shared task on sarcasm and sentiment detection in Arabic. In: Proceedings of the Sixth Arabic Natural Language Processing Workshop. Association for Computational Linguistics, pp 296–305. https://aclanthology.org/2021.wanlp-1.36',
    'Antoun W, Baly F, Hajj H (2020) AraBERT: transformer-based model for Arabic language understanding. In: Proceedings of the 4th Workshop on Open-Source Arabic Corpora and Processing Tools, with a Shared Task on Offensive Language Detection. European Language Resources Association, pp 9–15. https://aclanthology.org/2020.osact-1.2',
    'Belinkov Y, Bisk Y (2018) Synthetic and natural noise both break neural machine translation. In: International Conference on Learning Representations. https://arxiv.org/abs/1711.02173',
    'Conneau A, Khandelwal K, Goyal N, Chaudhary V, Wenzek G, Guzmán F, Grave E, Ott M, Zettlemoyer L, Stoyanov V (2020) Unsupervised cross-lingual representation learning at scale. In: Proceedings of the 58th Annual Meeting of the Association for Computational Linguistics. Association for Computational Linguistics, pp 8440–8451. https://doi.org/10.18653/v1/2020.acl-main.747',
    'Elmadany A, Mubarak H, Magdy W (2018) ArSAS: an Arabic speech-act and sentiment corpus of tweets. In: Proceedings of the 3rd Workshop on Open-Source Arabic Corpora and Processing Tools (OSACT 3), LREC 2018. European Language Resources Association. http://lrec-conf.org/workshops/lrec2018/W30/summaries/22_W30.html',
    'Gwet KL (2008) Computing inter-rater reliability and its variance in the presence of high agreement. Br J Math Stat Psychol 61(1):29–48. https://doi.org/10.1348/000711006X126600',
    'Hossain MM, Kovatchev V, Dutta P, Kao T, Wei E, Blanco E (2020) An analysis of natural language inference benchmarks through the lens of negation. In: Proceedings of the 2020 Conference on Empirical Methods in Natural Language Processing (EMNLP). Association for Computational Linguistics, pp 9106–9118. https://doi.org/10.18653/v1/2020.emnlp-main.732',
    'Inoue G, Alhafni B, Baimukan N, Bouamor H, Habash N (2021) The interplay of variant, size, and task type in Arabic pre-trained language models. In: Proceedings of the Sixth Arabic Natural Language Processing Workshop. Association for Computational Linguistics, pp 92–104. https://aclanthology.org/2021.wanlp-1.10',
    'Nabil M, Aly M, Atiya A (2015) ASTD: Arabic sentiment tweets dataset. In: Proceedings of the 2015 Conference on Empirical Methods in Natural Language Processing. Association for Computational Linguistics, pp 2515–2519. https://doi.org/10.18653/v1/D15-1299',
    'Pruthi D, Dhingra B, Lipton ZC (2019) Combating adversarial misspellings with robust word recognition. In: Proceedings of the 57th Annual Meeting of the Association for Computational Linguistics. Association for Computational Linguistics, pp 5582–5591. https://doi.org/10.18653/v1/P19-1561',
    'Ribeiro MT, Wu T, Guestrin C, Singh S (2020) Beyond accuracy: behavioral testing of NLP models with CheckList. In: Proceedings of the 58th Annual Meeting of the Association for Computational Linguistics. Association for Computational Linguistics, pp 4902–4912. https://doi.org/10.18653/v1/2020.acl-main.442',
    'Abu Farha I, Magdy W (2020) From Arabic sentiment analysis to sarcasm detection: the ArSarcasm dataset. In: Proceedings of the 4th Workshop on Open-Source Arabic Corpora and Processing Tools, with a Shared Task on Offensive Language Detection. European Language Resources Association, pp 32–39. https://aclanthology.org/2020.osact-1.5',
    'Abu Farha I, Magdy W (2021) A comparative study of effective approaches for Arabic sentiment analysis. Inf Process Manag 58(2):102438. https://doi.org/10.1016/j.ipm.2020.102438',
    'Badaro G, Baly R, Hajj H, El-Hajj W, Shaban KB, Habash N, Al-Sallab A, Hamdi A (2019) A survey of opinion mining in Arabic: a comprehensive system perspective covering challenges and advances in tools, resources, models, applications, and visualizations. ACM Trans Asian Low-Resour Lang Inf Process 18(3):1–52. https://doi.org/10.1145/3295662',
    'Clark JH, Garrette D, Turc I, Wieting J (2022) Canine: pre-training an efficient tokenization-free encoder for language representation. Trans Assoc Comput Linguist 10:73–91. https://doi.org/10.1162/tacl_a_00448',
    'Cohen J (1960) A coefficient of agreement for nominal scales. Educ Psychol Meas 20(1):37–46. https://doi.org/10.1177/001316446002000104',
    'Darwish K, Habash N, Abbas M, Al-Khalifa H, Al-Natsheh HT, Bouamor H, Bouzoubaa K, Cavalli-Sforza V, El-Beltagy SR, El-Hajj W, Jarrar M, Mubarak H (2021) A panoramic survey of natural language processing in the Arab world. Commun ACM 64(4):72–81. https://doi.org/10.1145/3447735',
    'Devlin J, Chang MW, Lee K, Toutanova K (2019) BERT: pre-training of deep bidirectional transformers for language understanding. In: Proceedings of the 2019 Conference of the North American Chapter of the Association for Computational Linguistics: Human Language Technologies, Volume 1 (Long and Short Papers). Association for Computational Linguistics, pp 4171–4186. https://doi.org/10.18653/v1/N19-1423',
    'Dror R, Baumer G, Shlomov S, Reichart R (2018) The hitchhiker’s guide to testing statistical significance in natural language processing. In: Proceedings of the 56th Annual Meeting of the Association for Computational Linguistics (Volume 1: Long Papers). Association for Computational Linguistics, pp 1383–1392. https://doi.org/10.18653/v1/P18-1128',
    'Ebrahimi J, Rao A, Lowd D, Dou D (2018) HotFlip: white-box adversarial examples for text classification. In: Proceedings of the 56th Annual Meeting of the Association for Computational Linguistics (Volume 2: Short Papers). Association for Computational Linguistics, pp 31–36. https://doi.org/10.18653/v1/P18-2006',
    'Efron B, Tibshirani RJ (1994) An introduction to the bootstrap. Chapman and Hall/CRC, New York. https://doi.org/10.1201/9780429246593',
    'Ettinger A (2020) What BERT is not: lessons from a new suite of psycholinguistic diagnostics for language models. Trans Assoc Comput Linguist 8:34–48. https://doi.org/10.1162/tacl_a_00298',
    'Gao J, Lanchantin J, Soffa ML, Qi Y (2018) Black-box generation of adversarial text sequences to evade deep learning classifiers. In: 2018 IEEE Security and Privacy Workshops (SPW). IEEE, pp 50–56. https://doi.org/10.1109/SPW.2018.00016',
    'Habash NY (2010) Introduction to Arabic natural language processing. Synthesis Lectures on Human Language Technologies. Morgan & Claypool. https://doi.org/10.2200/S00277ED1V01Y201008HLT010',
    'Jin D, Jin Z, Zhou JT, Szolovits P (2020) Is BERT really robust? A strong baseline for natural language attack on text classification and entailment. Proc AAAI Conf Artif Intell 34(05):8018–8025. https://doi.org/10.1609/aaai.v34i05.6311',
    'Kassner N, Schütze H (2020) Negated and misprimed probes for pretrained language models: birds can talk, but cannot fly. In: Proceedings of the 58th Annual Meeting of the Association for Computational Linguistics. Association for Computational Linguistics, pp 7811–7818. https://doi.org/10.18653/v1/2020.acl-main.698',
    'Kudo T, Richardson J (2018) SentencePiece: a simple and language independent subword tokenizer and detokenizer for neural text processing. In: Proceedings of the 2018 Conference on Empirical Methods in Natural Language Processing: System Demonstrations. Association for Computational Linguistics, pp 66–71. https://doi.org/10.18653/v1/D18-2012',
    'Kumar A, Makhija P, Gupta A (2020) Noisy text data: Achilles’ heel of BERT. In: Proceedings of the Sixth Workshop on Noisy User-generated Text (W-NUT 2020). Association for Computational Linguistics, pp 16–21. https://doi.org/10.18653/v1/2020.wnut-1.3',
    'Moradi M, Samwald M (2021) Evaluating the robustness of neural language models to input perturbations. In: Proceedings of the 2021 Conference on Empirical Methods in Natural Language Processing. Association for Computational Linguistics, pp 1558–1570. https://doi.org/10.18653/v1/2021.emnlp-main.117',
    'Obeid O, Zalmout N, Khalifa S, Taji D, Oudah M, Alhafni B, Inoue G, Eryani F, Erdmann A, Habash N (2020) CAMeL Tools: an open source Python toolkit for Arabic natural language processing. In: Proceedings of the Twelfth Language Resources and Evaluation Conference. European Language Resources Association, pp 7022–7032. https://aclanthology.org/2020.lrec-1.868',
    'Rosenthal S, Farra N, Nakov P (2017) SemEval-2017 task 4: sentiment analysis in Twitter. In: Proceedings of the 11th International Workshop on Semantic Evaluation (SemEval-2017). Association for Computational Linguistics, pp 502–518. https://doi.org/10.18653/v1/S17-2088',
    'Rust P, Pfeiffer J, Vulić I, Ruder S, Gurevych I (2021) How good is your tokenizer? On the monolingual performance of multilingual language models. In: Proceedings of the 59th Annual Meeting of the Association for Computational Linguistics and the 11th International Joint Conference on Natural Language Processing (Volume 1: Long Papers). Association for Computational Linguistics, pp 3118–3135. https://doi.org/10.18653/v1/2021.acl-long.243',
    'Schuster M, Nakajima K (2012) Japanese and Korean voice search. In: 2012 IEEE International Conference on Acoustics, Speech and Signal Processing (ICASSP). IEEE, pp 5149–5152. https://doi.org/10.1109/ICASSP.2012.6289079',
    'Sennrich R, Haddow B, Birch A (2016) Neural machine translation of rare words with subword units. In: Proceedings of the 54th Annual Meeting of the Association for Computational Linguistics (Volume 1: Long Papers). Association for Computational Linguistics, pp 1715–1725. https://doi.org/10.18653/v1/P16-1162',
    'van Miltenburg E, van der Lee C, Krahmer E (2021) Preregistering NLP research. In: Proceedings of the 2021 Conference of the North American Chapter of the Association for Computational Linguistics: Human Language Technologies. Association for Computational Linguistics, pp 613–623. https://doi.org/10.18653/v1/2021.naacl-main.51',
    'Wang X, Wang H, Yang D (2022) Measure and improve robustness in NLP models: a survey. In: Proceedings of the 2022 Conference of the North American Chapter of the Association for Computational Linguistics: Human Language Technologies. Association for Computational Linguistics, pp 4569–4586. https://doi.org/10.18653/v1/2022.naacl-main.339',
    'Wiegand M, Balahur A, Roth B, Klakow D, Montoyo A (2010) A survey on the role of negation in sentiment analysis. In: Proceedings of the Workshop on Negation and Speculation in Natural Language Processing. University of Antwerp, pp 60–68. https://aclanthology.org/W10-3111',
    'Wolf T, Debut L, Sanh V, Chaumond J, Delangue C, Moi A, Cistac P, Rault T, Louf R, Funtowicz M, Davison J, Shleifer S, von Platen P, Ma C, Jernite Y, Plu J, Xu C, Le Scao T, Gugger S, Drame M, Lhoest Q, Rush A (2020) Transformers: state-of-the-art natural language processing. In: Proceedings of the 2020 Conference on Empirical Methods in Natural Language Processing: System Demonstrations. Association for Computational Linguistics, pp 38–45. https://doi.org/10.18653/v1/2020.emnlp-demos.6',
    'Xue L, Barua A, Constant N, Al-Rfou R, Narang S, Kale M, Roberts A, Raffel C (2022) ByT5: towards a token-free future with pre-trained byte-to-byte models. Trans Assoc Comput Linguist 10:291–306. https://doi.org/10.1162/tacl_a_00461',
]
REFERENCES.sort(key=str.casefold)

ABSTRACT = (
    'Arabic tweets contain orthographic variation that readers barely notice: the elongation character tatweel, '
    'emphatic letter repetition, and inconsistent spelling of hamza and final letters. We measure how this variation '
    'affects five Arabic and multilingual encoders fine-tuned for three-way sentiment classification on two tweet '
    'corpora, ArSarcasm-v2 and ArSAS, with three training seeds per model. The study was preregistered and frozen '
    'before a single evaluation on the held-out test sets. At the input level the effect is stark: in all four Arabic '
    'WordPiece models, one tatweel turns the edited word into an unknown token, including every edited negation word. '
    'Predictions are sturdier. On the test sets, tatweel in one word changes 3–9% of predictions, and tatweel in 30% '
    'of the words changes 10–22%. Tokenizer damage predicts these rates only weakly, but it matters for what the '
    'changes cost: under dense tatweel the Arabic models lose 3.7–5.6 Macro-F1 points, against 1.4–2.7 for XLM-R. '
    'Two annotators judged every sampled tatweel edit to preserve meaning and sentiment, so these changes are errors; '
    'elongation, by contrast, was read as a change of intensity. Editing the negation word itself adds only a small '
    'and inconsistent effect. Removing tatweel and collapsing letter runs before tokenization undoes the edits at no '
    'measurable cost on clean text, and on test tweets that already contain letter runs it raises Macro-F1 by 4–8 '
    'points on ArSarcasm-v2. A negation-aware normalizer adds nothing. We release the perturbation suite, '
    'leakage-safe splits, predictions and the preregistration.')
KEYWORDS = ['Arabic sentiment analysis', 'Robustness', 'Tokenization', 'Orthographic noise', 'Negation',
            'Preregistration']


def manuscript(data, figures, output, anonymous=True):
    doc = base_document()
    add_line_numbers(doc.sections[0])
    para(doc, TITLE[0].upper() + TITLE[1:], align='center', size=14, bold=True, after=10)
    if not anonymous:
        para(doc, '[[Author names and affiliations: see title page]]', align='center')
    heading(doc, 'Abstract', 1)
    para(doc, ABSTRACT)
    para(doc, '**Keywords** ' + ' · '.join(KEYWORDS), align='left', after=8)

    heading(doc, '1 Introduction', 1)
    para(doc, 'Social media Arabic is written with considerable freedom. Writers stretch a word with the tatweel '
              'character for visual effect (جميـــل for جميل “beautiful”), repeat a letter for emphasis (حلووو, '
              'كثيييير), and often omit hamza or write final ة as ه. In the two corpora we study, 3–5% of training and '
              'validation tweets contain a tatweel, and a similar share repeat a letter three or more times (Table 7). '
              'Native readers glide over these variants. Whether sentiment classifiers do so is less clear, because '
              'the tokenizers that feed them were built from vocabularies in which these forms are rare or absent.')
    para(doc, 'We arrived at this question through a negative result. A negation-aware normalization method we '
              'developed earlier did no better than a simpler normalizer under any condition. In that experiment, '
              'single-word edits applied to one MARBERT model produced accuracy changes too small to separate from '
              'zero. We concluded that the design, not the models, made the experiment uninformative. A single model '
              'gives no basis for comparison, a single edited word per tweet is a weak treatment, and the repetition '
              'edits skipped the long vowels ا, و and ي, which carry most emphatic repetition in real tweets.')
    para(doc, 'This paper replaces that experiment with a broader, preregistered audit. We fine-tune five encoders '
              'with three seeds each on two tweet datasets. We then apply four families of orthographic edits at two '
              'densities, and in some conditions we edit the negation word itself. Two annotators, and a third for '
              'their disagreements, judged whether the edits preserve meaning and sentiment. After development on the '
              'validation sets, we froze the study and evaluated it once on the test sets. We ask four questions: '
              '(RQ1) how often do realistic edits change predictions; (RQ2) do tokenizer effects explain differences '
              'between models; (RQ3) is editing the negation word worse than editing another word; and (RQ4) does '
              'simple normalization before tokenization undo the damage?')
    para(doc, 'The main findings, all of which hold on the held-out test sets, are:')
    bullets(doc, [
        'Arabic WordPiece tokenizers map any word containing tatweel to [UNK]. This includes negation words, so an '
        'edited tweet such as لـا reaches the model with its negation removed (Section 4).',
        'Predictions change far less than this input damage would suggest. Tatweel in one word changes 3–9% of them, '
        'and tatweel in 30% of the words changes 10–22%. The Arabic models’ changes are costlier, however: on '
        'ArSarcasm-v2 each point of flip rate costs them about 2.5 times as much Macro-F1 as it costs XLM-R '
        '(Section 6).',
        'Editing the negation word adds at most 2.5 points of flip rate over editing another word, and the difference '
        'is significant in only half of the test comparisons (Section 7).',
        'Stripping tatweel and collapsing letter runs removes nearly all of these changes at no measurable cost on '
        'clean text. On test tweets with natural letter runs it also improves accuracy. The negation guard of our '
        'earlier method adds nothing (Sections 8 and 9).'])

    heading(doc, '2 Related work', 1)
    heading(doc, '2.1 Arabic sentiment analysis', 2)
    para(doc, 'Arabic sentiment analysis has moved from lexicon- and feature-based systems to fine-tuned '
              'transformers, a development traced by several surveys (Badaro et al. 2019; Darwish et al. 2021). '
              'Three-way sentiment on tweets became a standard task through SemEval, which included an Arabic '
              'subtask (Rosenthal et al. 2017), and through dedicated corpora such as ASTD (Nabil et al. 2015), ArSAS '
              '(Elmadany et al. 2018) and ArSarcasm (Abu Farha and Magdy 2020), later extended to ArSarcasm-v2 for '
              'the WANLP 2021 shared task (Abu Farha et al. 2021). Comparative studies find that fine-tuned BERT-style '
              'models outperform earlier approaches on these datasets (Abu Farha and Magdy 2021). This literature '
              'reports accuracy on the text as distributed. It rarely asks how stable a model is when the same tweet '
              'is written differently, although written Arabic tolerates a great deal of such variation: tatweel '
              'is a purely typographic device, hamza is often omitted, and final ya and ta marbuta are written in '
              'several ways (Habash 2010). Toolkits such as CAMeL Tools (Obeid et al. 2020) therefore provide '
              'orthographic normalization, but whether a given model needs it is usually left to the user.')
    heading(doc, '2.2 Subword tokenization and Arabic pretrained models', 2)
    para(doc, 'BERT (Devlin et al. 2019) and its Arabic descendants segment text with WordPiece (Schuster and '
              'Nakajima 2012), a greedy subword algorithm related to byte-pair encoding (Sennrich et al. 2016). A '
              'word that contains a character outside the vocabulary cannot be segmented and is replaced by a single '
              'unknown token. MARBERT (Abdul-Mageed et al. 2021), AraBERT (Antoun et al. 2020) and CAMeLBERT (Inoue '
              'et al. 2021) use WordPiece vocabularies learned from Arabic text; MARBERT and its successor share a '
              'vocabulary built from tweets. XLM-R (Conneau et al. 2020) instead uses a multilingual SentencePiece '
              'model (Kudo and Richardson 2018) that falls back on individual characters and therefore never '
              'produces an unknown token. Tokenizer quality has been shown to explain part of the gap between '
              'monolingual and multilingual models, Arabic included (Rust et al. 2021), and tokenization-free '
              'encoders have been proposed to remove the vocabulary bottleneck altogether (Clark et al. 2022; Xue et '
              'al. 2022). The recommended preprocessing for some Arabic models removes tatweel, but nothing in the '
              'model interface enforces this. We are not aware of work that measures what omitting that step costs '
              'a fine-tuned model.')
    heading(doc, '2.3 Robustness to noise', 2)
    para(doc, 'Both synthetic and natural character noise degrade neural models (Belinkov and Bisk 2018). '
              'Adversarial work showed that a few character flips or word substitutions can change a classifier’s '
              'decision (Ebrahimi et al. 2018; Gao et al. 2018; Jin et al. 2020), and non-adversarial studies found '
              'that pretrained models lose accuracy under realistic spelling noise (Kumar et al. 2020; Moradi and '
              'Samwald 2021). Character-level recognizers have been proposed to defend against misspellings (Pruthi '
              'et al. 2019), and behavioural test suites probe whether a model is invariant to such edits (Ribeiro et '
              'al. 2020); Wang et al. (2022) survey the field. Most of this work concerns English and treats every '
              'perturbation as label-preserving. Our edits are specific to Arabic script. We separate edits presumed '
              'to preserve meaning from a typo control and verify that separation with human judgments, and we '
              'report the costs of a defense alongside its benefits.')
    heading(doc, '2.4 Negation', 2)
    para(doc, 'Negation is a classic difficulty in sentiment analysis (Wiegand et al. 2010). Pretrained language '
              'models handle it poorly in diagnostic probes (Ettinger 2020; Kassner and Schütze 2020), and benchmarks '
              'under-represent it, so models trained on them often ignore it (Hossain et al. 2020). We do not parse '
              'negation scope. Instead, we compare an edit on the negation word with the same edit on another word '
              'in the same tweet, which isolates the contribution of the negation word itself.')

    heading(doc, '3 Study design', 1)
    heading(doc, '3.1 Protocol', 2)
    para(doc, 'Following the case for preregistration in NLP (van Miltenburg et al. 2021), we preregistered the '
              'conditions, metrics, hypotheses and decision rules before observing any model '
              'prediction on the new edits, and we developed the analysis on the validation sets. A written amendment '
              'fixed the test protocol. The study was then frozen at a tagged commit, and only after that were the '
              'test sets perturbed and evaluated, once. We report the test sets as the main results; the validation '
              'results are given in Online Resource 1. The tokenizer audit in Section 4 was run before the '
              'preregistration and is disclosed there.')
    heading(doc, '3.2 Data', 2)
    para(doc, 'ArSarcasm-v2 (Abu Farha et al. 2021) contains dialectal and MSA tweets labelled negative, neutral or '
              'positive. We use a leakage-safe split that removes from training any tweet whose normalized text '
              'occurs in the official test set. This gives 10,037 training, 2,509 validation and 3,000 official test '
              'tweets. ArSAS (Elmadany et al. 2018) contains 21,064 tweets with an additional Mixed class, which we '
              'drop. We merge repeated tweet IDs and identical texts into groups, then exclude the 90 tweets whose '
              'groups carry conflicting labels and 44 further copies linked to excluded tweets. A stratified group '
              'split then gives 13,738 training, 1,966 validation and 3,923 test tweets. We considered ASTD (Nabil et '
              'al. 2015) as the second dataset and rejected it, because 20% of ArSarcasm-v2 training tweets (2,473) '
              'also occur in ASTD. ArSAS shares one tweet with ArSarcasm-v2, and that tweet is excluded. Each model '
              'is trained and evaluated separately on each dataset.')
    heading(doc, '3.3 Edits', 2)
    para(doc, 'Table 1 illustrates the four families. *Tatweel* inserts one ـ after the last letter that joins to '
              'the next. *Elongation* repeats the last internal long vowel, or the final letter if there is none, two '
              'or four extra times. *Folding* writes أ, إ and آ as ا, final ة as ه, and final ى as ي. *Typo* replaces '
              'one interior letter with its neighbour on the Arabic keyboard. The typo family was intended as a '
              'control that damages meaning; the human judgments in Section 5 show that it does so only partly.')
    table(doc, 1, 'The edit families, illustrated on single words',
          ['Family', 'Original', 'Edited'],
          [['Tatweel', 'حلو', 'حلـو'], ['Elongation ×2', 'حلو', 'حلووو'], ['Elongation ×4', 'كثير', 'كثيييير'],
           ['Folding', 'أحب / جميلة', 'احب / جميله'], ['Typo (control)', 'جيد', 'جبد'], ['Cue (tatweel)', 'لا', 'لـا']],
          widths=[4.0, 3.0, 3.0], size=9)
    para(doc, 'Each family is applied at two densities: *single* edits one word, and *dense* edits 30% of the '
              'Arabic words in the tweet, rounded up. In *cue* conditions the edit falls on an unambiguous negation word (لا, لم, '
              'لن, ليس, مش, مو, مافي, بدون and their variants). We leave out ما, whose relative, interrogative and '
              'exclamative uses outnumber its negating use. Non-cue conditions never edit a negation word. They also '
              'skip URLs, mentions, hashtags and tokens attached to Latin characters. Which word gets edited is '
              'decided by hashing the tweet ID and the condition. Single edits are a subset of the dense edits for '
              'the same tweet, and both elongation levels edit the same character.')
    para(doc, 'On the test sets, edits apply to between 2,119 and 2,950 of the 3,000 ArSarcasm-v2 tweets per '
              'condition, and to between 3,303 and 3,898 of the 3,923 ArSAS tweets. Cue conditions apply to the 626 '
              'and 623 test tweets that contain an editable negation word. Across both partitions this amounts to '
              'about 100,000 edited tweets per model.')
    heading(doc, '3.4 Models', 2)
    para(doc, 'We fine-tune MARBERT, MARBERTv2, AraBERTv02-twitter, CAMeLBERT-DA and XLM-R-base with one fixed '
              'recipe in the Transformers library (Wolf et al. 2020): AdamW with a learning rate of 2×10⁻⁵, batch size 16 and at most five epochs; early stopping '
              'and checkpoint selection on validation Macro-F1; and seeds 42, 43 and 44. This gives 30 models. '
              'Revisions of all pretrained weights are pinned. Table 2 reports clean performance. The official '
              'ArSarcasm-v2 test set is markedly harder than its validation split, and every model scores 8–12 '
              'points lower on it.')
    groups, header, rows = clean_table(data)
    table(doc, 2, 'Clean Macro-F1 (%). Validation: mean ± s.d. over three seeds. Test: seed-averaged, with 95% '
                  'cluster-bootstrap interval, on the 2,948 and 3,898 test tweets eligible for single-word tatweel',
          header, rows, groups=groups, widths=[2.4, 2.4, 2.4, 3.4, 3.4])
    heading(doc, '3.5 Defenses', 2)
    para(doc, 'At inference time we compare preprocessors, each applied identically to clean and edited text. RAW '
              'leaves the input unchanged. STD removes tatweel and collapses runs of three or more identical letters '
              'to a single letter. NO_CTX and NASN2 are the two conservative normalizers from our earlier work; they '
              'share all of their rules except one, in that NASN2 also declines to collapse repeated letters near a '
              'negation word. Their repetition rule covers consonants only, so they do not undo elongated long '
              'vowels. The comparison between these two was settled on validation (H5 below), and the test protocol '
              'therefore evaluates RAW and STD only.')
    heading(doc, '3.6 Metrics and statistics', 2)
    para(doc, 'The primary metric is the *flip rate*, the share of edited tweets whose predicted class differs from '
              'the prediction for the clean tweet. It needs no gold label and isolates the effect of the edit. We '
              'also report Macro-F1 on clean and edited versions of the same tweets and the difference between the '
              'two.')
    para(doc, 'Intervals are 95% percentiles of 2,000 cluster-bootstrap resamples (Efron and Tibshirani 1994; Dror et '
              'al. 2018). Clusters are groups of duplicate '
              'texts, and the rate for each resample is averaged over the three seeds. Contrasts are paired on '
              'identical tweets. The preregistered hypotheses are: H1, dense edits change more predictions than '
              'single edits for the three presumed-preserving families; H2, under dense tatweel each WordPiece model '
              'flips more often than XLM-R; H3, tatweel on the negation word changes more predictions than tatweel '
              'on another word of the same tweet; H4, STD reduces flip rates relative to RAW; and H5, NASN2 reduces '
              'flip rates relative to NO_CTX. A hypothesis counts as supported for a model when its interval '
              'excludes zero in the stated direction.')
    heading(doc, '3.7 Deviation from the preregistration', 2)
    para(doc, 'The preregistered zero-shot large language model (LLM) was not evaluated, because the required GPU '
              'was unavailable.')
    heading(doc, '3.8 Use of AI tools', 2)
    para(doc, '[[Authors: confirm or edit this paragraph before submission.]] A large language model (Claude, '
              'Anthropic) was used to assist in writing analysis code and in drafting and editing the manuscript. '
              'The study design, all analyses and every reported number were checked by the authors, who take full '
              'responsibility for the content.')

    heading(doc, '4 What the tokenizer sees', 1)
    figure(doc, 1, figures[1], 'Number of tokens produced for an edited word, relative to the original word, '
           'averaged over single-word edits on the validation sets. A hollow marker means the edited word becomes '
           '[UNK] (rate 1.00 in every such cell). The test sets give the same pattern. MARBERT and MARBERTv2 share '
           'a vocabulary, so their rows coincide')
    para(doc, 'Figure 1 shows, for every model and edit family, how many tokens an edited word receives relative to '
              'the original. The pattern is the same on both datasets and both partitions.')
    para(doc, 'None of the four Arabic WordPiece vocabularies contains the tatweel character. WordPiece therefore '
              'replaces any word containing it with a single [UNK]. This happens in 100% of tatweel edits, including '
              'every edited negation word (1,969 across the four partitions). A tweet containing لـا thus reaches '
              'these models with no negation. XLM-R never produces an unknown token. It instead splits a word '
              'containing tatweel into 2.3–2.6 times as many pieces, and an edited negation word into 2.9 times as '
              'many. Elongation fragments the word in every model: four extra letters multiply its token count by '
              '2.4–4.1. Folding is largely absorbed by MARBERT, whose lowercasing strips hamza from alef. As a '
              'result, 31–39% of its single-word folding edits leave the token sequence unchanged. The other models '
              'see nearly every folding edit.')

    heading(doc, '5 Do humans notice the edits?', 1)
    para(doc, 'The study rests on the claim that most of these edits preserve meaning. To test it, two native-speaker '
              'annotators independently judged 330 edited tweets drawn from the ArSarcasm-v2 validation set, 30 per '
              'condition and one edit per tweet. They answered five questions per item. Their first 20 items served '
              'as a pilot and are excluded from the agreement figures. Agreement (Table 3) is high on every field. '
              'Cohen’s κ (Cohen 1960) is low for meaning preservation because almost every answer is “yes”. Gwet’s AC1 (Gwet '
              '2008) corrects for this skew and gives 0.88.')
    header, rows = agreement_table(data)
    table(doc, 3, 'Agreement between the two annotators on 310 items (20 pilot items excluded)', header, rows,
          widths=[5.0, 2.4, 2.4, 2.4], size=9)
    para(doc, 'A third annotator resolved the 87 items on which the two disagreed about meaning or sentiment '
              'preservation. That annotator saw both earlier answers, without the annotators’ identities or notes and '
              'in random order. On meaning, the third annotator sided with each of the two about equally often (76% '
              'and 75%). On sentiment, the third annotator sided more often with annotator A (52%) than with '
              'annotator B (21%), who had marked many of these items unsure.')
    header, rows = preservation_table(data)
    table(doc, 4, 'Items, out of 30 per condition, judged meaning-preserving by annotator A and by annotator B; '
                  'items for which both judged meaning preserved, or both meaning and sentiment preserved; and the '
                  'final count of meaning- and sentiment-preserving items after a third annotator resolved the '
                  'disagreements', header, rows, widths=[5.4, 2.6, 2.4, 3.4, 1.6])
    para(doc, 'The results (Table 4) separate the families clearly.')
    bullets(doc, [
        '*Tatweel* preserved both meaning and sentiment in all 90 items, including the 30 in which it fell on a '
        'negation word. Any prediction change under tatweel is therefore an error.',
        '*Elongation* preserved meaning in all 120 items, but it often changed sentiment. The sentiment was judged '
        'unchanged in only 17 and 14 of 30 items for one-word elongation by two and four letters, in 6 of 30 when '
        '30% of the words were elongated, and in none of the 30 items with an elongated negation word. Readers take '
        'repetition as emphasis, so a prediction change under elongation is not necessarily an error.',
        '*Folding* preserved meaning in 48 of 60 items by agreement, and in all 60 after adjudication. The disputed '
        'items wrote على as علي or dropped a hamza, which the third annotator judged cosmetic.',
        '*Typos*, meant as the control, were often read through. After adjudication, 19 of 30 single-word typos and '
        '8 of 30 dense ones preserved meaning and sentiment.'])
    para(doc, 'Adjudication changed none of the counts for tatweel or elongation. Every elongation item it resolved '
              'was judged meaning-preserving but sentiment-changing.')

    heading(doc, '6 How often do predictions change?', 1)
    figure(doc, 2, figures[2], 'Test-set flip rates with 95% cluster-bootstrap intervals (RAW inputs, three seeds '
           'averaged per model). Full tables for both partitions are in Online Resource 1')
    para(doc, 'Figure 2 shows the test flip rates. A single edited word rarely matters. Tatweel in one word changes '
              '7.0–9.3% of predictions on ArSarcasm-v2 and 3.0–3.9% on ArSAS; two extra letters change 2.9–7.8%; '
              'and folding changes 0.9–4.3%.')
    para(doc, 'Density matters a great deal. When 30% of the words carry tatweel, flip rates rise to 13.4–21.6% on '
              'ArSarcasm-v2 and 9.6–11.3% on ArSAS. Dense elongation reaches 16.2% for XLM-R and 15.3% for '
              'CAMeLBERT-DA. Every dense-versus-single contrast for the three presumed-preserving families is '
              'positive, in all five models and on both datasets (30 of 30 on test, as on validation; H1). The '
              'increases range from 1.1 to 3.8 points for folding and from 6.1 to 12.2 points for tatweel and '
              'elongation. Tatweel changes more predictions on the ArSarcasm-v2 test set than on its validation '
              'split, in line with the lower clean accuracy, while elongation and typos change slightly fewer.')
    heading(doc, '6.1 Tokenizer damage barely predicts flip rates', 2)
    para(doc, 'The tokenizer audit predicts a large gap between WordPiece models, which lose the edited word '
              'entirely, and XLM-R, which keeps a fragmented version. The gap exists but depends on the dataset. '
              'Under dense tatweel on ArSarcasm-v2, every WordPiece model flips more often than XLM-R, by 3.8–8.1 '
              'points. On ArSAS three of the four do so, by 1.4–1.8 points (H2: 7 of 8 test contrasts, 6 of 8 on '
              'validation). On ArSarcasm-v2, XLM-R is the most sensitive model to elongation.')
    para(doc, 'Across the 55 model–condition cells per dataset, the rank correlation between the [UNK] rate of the '
              'edited words and the flip rate is 0.44 on the ArSarcasm-v2 test set but only 0.07 on the ArSAS test '
              'set. On validation it is 0.18 and 0.14. Fragmentation does not correlate with flip rate on any '
              'partition (|ρ| ≤ 0.11). One explanation is that fine-tuning already exposes the models to [UNK] for '
              'tatweel words, since 4.6% of ArSarcasm-v2 training and validation tweets contain a natural tatweel.')
    heading(doc, '6.2 WordPiece errors cost more', 2)
    para(doc, 'The flip rate does not say whether a changed prediction is wrong; Table 5 does. Under dense tatweel on '
              'the test sets, Macro-F1 falls by 4.9–5.6 points for the WordPiece models on ArSarcasm-v2 and by '
              '3.7–4.0 points on ArSAS. For XLM-R it falls by 1.4 and 2.7 points. Per point of flip rate, the '
              'WordPiece models lose 0.25–0.28 Macro-F1 points on ArSarcasm-v2, against 0.10 for XLM-R. On ArSAS '
              'the gap is smaller: 0.33–0.39 against 0.28. Validation shows the same ordering. This is consistent '
              'with an erased sentiment word removing the evidence for the label rather than merely adding noise. '
              'Dense typos cost the most on ArSAS (7.1–11.4 points).')
    groups, header, rows = drop_table(data.test, data)
    table(doc, 5, 'Test-set Macro-F1 drop in points (clean minus edited, same tweets) under dense edits, with 95% '
                  'intervals', header, rows, groups=groups, size=7.5, widths=[2.2] + [2.35] * 6)

    heading(doc, '7 Negation as a stress test', 1)
    para(doc, 'Natural noise on a negation word is rare. In the training and validation tweets, at most five per '
              'corpus contain an elongated or tatweel-marked negation word. We therefore treat the cue conditions as '
              'a controlled probe of how models represent negation, not as an estimate of real-world risk.')
    para(doc, 'For each tweet with an editable negation word, we compare editing that word with editing another word '
              'in the same tweet (H3). Tatweel on the negation word, which turns it into [UNK], always changes at '
              'least as many predictions as tatweel elsewhere. The difference is small, however. On the test sets '
              'the increase is 0.3–2.5 points, and it is significant in four of the eight WordPiece model–dataset '
              'pairs (AraBERT on ArSarcasm-v2; MARBERT, MARBERTv2 and CAMeLBERT-DA on ArSAS). On validation it was '
              'significant in seven of eight, with increases of 2.0–3.3 points. Across both partitions, losing the '
              'negation word changes the predicted class of at most one additional tweet in thirty. These models '
              'evidently do not rest their decisions on the negation word alone.')
    para(doc, 'Elongating the negation word (لاااا) increases flip rates significantly in only three of ten test '
              'pairs, and in two of ten on validation. The annotators read this edit very differently from tatweel. '
              'Annotator A judged the sentiment changed in 26 of 30 items, and annotator B in 16 (with 13 marked '
              'unsure). After adjudication, none of the 30 was judged to keep its sentiment (Section 5). An '
              'elongated لاااا is an emphatic refusal, and a model that changes its prediction may be reading it '
              'correctly.')

    heading(doc, '8 Does normalization help?', 1)
    para(doc, 'Table 6 compares STD with RAW on the test sets. STD removes tatweel and collapses letter runs, and so '
              'reverses exactly the edits it targets. It brings flip rates under tatweel to zero and under dense '
              'elongation to at most 0.4% (H4: 70 of 70 tests on test, as on validation). It does not touch folding '
              'or typos.')
    groups, header, rows = defense_table(data.test, data, ['RAW', 'STD'])
    table(doc, 6, 'Test-set flip rates (%) under dense edits, without (RAW) and with STD normalization of clean '
                  'and edited text', header, rows, groups=groups, widths=[2.8, 2.6, 2.0, 2.0, 2.0, 2.0], size=8.5, left=2)
    para(doc, 'The more relevant question is what STD costs on unedited text, where natural letter runs may carry '
              'emphasis. The cost is negligible. On the clean test tweets used in the tatweel conditions (2,948 and '
              '3,898), Macro-F1 under STD differs from RAW by at most 0.35 points for every model, and it is higher '
              'under STD for six of the ten model–dataset pairs.')
    para(doc, 'On validation we also compared the two conservative normalizers from our earlier work (Online '
              'Resource 1, Table S5). They remove most tatweel effects, leaving 0.5–1.0% of flips because they keep '
              'tatweel inside the words they treat as negation markers. They do not undo elongated long vowels, '
              'which their repetition rule excludes. The negation guard that distinguishes NASN2 from NO_CTX changes '
              'nothing that matters. Across 110 paired contrasts the two differ by at most 0.2 points. None favours '
              'NASN2, and five favour NO_CTX by fractions of a point. Under the preregistered rule, the claim that '
              'negation-aware normalization improves robustness is not supported (H5).')

    heading(doc, '9 Natural noise', 1)
    para(doc, 'Controlled edits show what can happen; natural noise shows what does (Table 7). The ArSarcasm-v2 test '
              'set is noisier still: 219 of its 3,000 tweets (7.3%) contain a tatweel and 339 contain a letter run. '
              'For these tweets we compare the prediction for the tweet as written with the prediction after STD '
              'cleaning (Table 8).')
    table(doc, 7, 'Naturally occurring noise in the training and validation partitions',
          ['', 'ArSarcasm-v2', 'ArSAS'],
          [['Tweets (train + val.)', '12,546', '15,704'], ['with tatweel', '571 (4.6%)', '502 (3.2%)'],
           ['with a letter run ≥3', '498 (4.0%)', '436 (2.8%)'], ['   of which a long vowel', '300 (2.4%)', '294 (1.9%)'],
           ['tatweel on a negation word', '1', '1'], ['letter run on a negation word', '5', '3']],
          widths=[5.4, 3.0, 3.0], size=9)
    groups, header, rows = natural_table(data)
    table(doc, 8, 'Test tweets that already contain noise: flip rate (%) between the tweet as written and its '
                  'STD-cleaned form, with 95% intervals', header, rows, groups=groups, size=8,
          widths=[2.4, 3.4, 3.4, 3.4, 3.4])
    para(doc, 'Cleaning natural noise moves many predictions. On the ArSarcasm-v2 test set, removing tatweel changes '
              '7–12% of them and collapsing letter runs 16–20%. On the ArSAS test set the figures are 1–4% and 4–9%.')
    para(doc, 'For letter runs on ArSarcasm-v2 the change is also an improvement. Macro-F1 is higher on the cleaned '
              'tweets for all five models, by 3.9 to 7.8 points, and every interval lies above zero. For natural '
              'tatweel on ArSarcasm-v2, on the smaller ArSAS slices (n = 89 and 128) and on the validation slices '
              '(n = 42–123), the differences are not significant. Collapsing a run removes emphasis that human '
              'readers notice (Section 5). For these classifiers, however, the emphasized form appears to be harder '
              'to read than the plain word.')

    heading(doc, '10 Discussion', 1)
    para(doc, 'Table 9 summarizes the preregistered hypotheses on both partitions. The tokenizer view and the '
              'prediction view tell different stories. A WordPiece model receives an [UNK] in place of every tatweel '
              'word, which sounds catastrophic. Yet the prediction usually survives, because sentiment in a tweet '
              'rarely depends on a single word. Where the two views agree is in the cost of the predictions that do '
              'change. When an Arabic WordPiece model loses a word and changes its prediction, the change costs more '
              'Macro-F1 than a change made by XLM-R, which merely fragments the word. On ArSarcasm-v2 it costs about '
              '2.5 times as much.')
    groups, header, rows = hypotheses_table(data)
    table(doc, 9, 'Supported tests per preregistered hypothesis (95% interval excludes zero in the stated '
                  'direction). H5 compares two normalizers that the test protocol did not include', header, rows,
          groups=groups, widths=[2.4, 2.6, 2.6, 2.6, 2.6], size=9)
    para(doc, 'The families also call for different readings. Tatweel is pure typography. Annotators saw no change '
              'in 90 of 90 items, so every prediction change under tatweel is an error, and a trivial fix removes '
              'it. Elongation is not noise in the same sense: readers took it as a change of intensity in most '
              'items, and a model that responds to it is not necessarily wrong. Benchmarks that treat all '
              'orthographic variation as a meaning-preserving perturbation conflate these cases.')
    para(doc, 'The negation results are a caution against over-interpreting tokenizer audits. Deleting the negation '
              'word from a WordPiece model’s input sounds as if it should invert many predictions. In practice it '
              'adds only a few, and the effect did not fully replicate on the test sets. The models appear to read '
              'sentiment from the rest of the tweet as well.')
    para(doc, 'The practical recommendation is straightforward. Anyone deploying an Arabic WordPiece model on social '
              'media text should remove tatweel before tokenization. The cost on clean text stayed within 0.35 '
              'Macro-F1 points for every model, and on ArSarcasm-v2 cleaning natural letter runs improved accuracy '
              'as well. Collapsing letter runs does remove emphasis that readers perceive, so it should be a '
              'deliberate choice. Negation-aware exceptions to these rules did not help in our experiments.')
    heading(doc, '10.1 Limitations', 2)
    para(doc, 'The edits are rule-based and applied one family at a time, whereas real noise co-occurs. The negation '
              'lexicon is finite and leaves out ما. The human validation covers one dataset and 30 items per '
              'condition. The natural-noise slices are small, and their accuracy effects are significant only for '
              'letter runs on ArSarcasm-v2. Both datasets consist of tweets, and other genres may behave '
              'differently. We did not evaluate instruction-tuned LLMs, whose byte-level tokenizers resemble '
              'XLM-R’s but whose behaviour under these edits remains to be measured, as does that of '
              'tokenization-free encoders (Clark et al. 2022; Xue et al. 2022).')

    heading(doc, '11 Conclusion', 1)
    para(doc, 'Orthographic variation that Arabic readers ignore reaches Arabic WordPiece models as missing words: a '
              'single tatweel erases the word it touches, negation words included. In a preregistered audit of 30 '
              'fine-tuned models, the resulting prediction changes were smaller than this input damage suggests and '
              'only weakly predicted by it, but each change cost the Arabic models more accuracy than it cost '
              'XLM-R. Removing tatweel and collapsing letter runs before tokenization undid the controlled edits at '
              'no measurable cost and improved accuracy on naturally elongated tweets, whereas a negation-aware '
              'refinement of that rule added nothing. Human judgments showed that tatweel is purely typographic '
              'while elongation carries intensity, which robustness benchmarks for Arabic should keep apart.')

    heading(doc, 'Supplementary Information', 1)
    para(doc, 'Online Resource 1 contains the flip rate for every condition and model on both partitions, the '
              'validation results for all four defenses and the validation Macro-F1 drops.')
    heading(doc, 'Statements and Declarations', 1)
    for name, text in [
            ('Funding', '[[Authors: state funding, or: The authors declare that no funds, grants, or other support '
                        'were received during the preparation of this manuscript.]]'),
            ('Competing interests', '[[Authors: confirm.]] The authors have no relevant financial or non-financial '
                                    'interests to disclose.'),
            ('Ethics approval', 'Not applicable. The study uses publicly available research datasets; only tweet IDs, '
                                'hashes, labels and aggregate results are released.'),
            ('Author contributions', '[[Authors: describe each author’s contribution (CRediT roles) after review.]]'),
            ('Data availability', 'Code, configurations, the preregistration and all text-free artefacts (perturbation '
                                  'logs, split manifests, predictions keyed by text hash and bootstrap results) are '
                                  'available at [[repository link withheld for double-blind review]]. Tweet texts are '
                                  'available from the providers of ArSarcasm-v2 and ArSAS under their licenses. Every '
                                  'table in this article is generated from the released analysis bundles by a script.')]:
        para(doc, f'**{name}** {text}')

    heading(doc, 'References', 1)
    for reference in REFERENCES:
        paragraph = para(doc, reference, align='left', size=9, after=3)
        paragraph.paragraph_format.left_indent = Cm(0.5)
        paragraph.paragraph_format.first_line_indent = Cm(-0.5)
    doc.save(output)


def supplementary(data, output):
    doc = base_document()
    para(doc, '**Online Resource 1**', align='left', size=12)
    para(doc, TITLE[0].upper() + TITLE[1:], align='left', italic=True, after=10)
    number = 1
    for bundle, partition in [(data.test, 'test'), (data.validation, 'validation')]:
        for dataset, name in DATASETS.items():
            header, rows = flip_table(bundle, data, dataset)
            table(doc, f'S{number}', f'{name} {partition}: flip rate (%) for every condition, RAW inputs', header,
                  rows, widths=[5.4] + [2.0] * 5, size=8.5)
            number += 1
    groups, header, rows = defense_table(data.validation, data, ['RAW', 'STD', 'NO_CONTEXT', 'NASN2'])
    table(doc, 'S5', 'Validation flip rates (%) under dense edits for all four preprocessing defenses, including the '
                     'conservative normalizers NO_CTX and NASN2', header, rows, groups=groups, size=8,
          widths=[2.0, 2.3] + [1.48] * 8, left=2)
    groups, header, rows = drop_table(data.validation, data)
    table(doc, 'S6', 'Validation Macro-F1 drop in points under dense edits, with 95% intervals', header, rows,
          groups=groups, size=7.5, widths=[2.2] + [2.35] * 6)
    doc.save(output)


def title_page(output):
    doc = base_document()
    para(doc, '**Title page**', align='left', size=12, after=10)
    para(doc, TITLE[0].upper() + TITLE[1:], align='left', size=14, bold=True, after=12)
    para(doc, '[[First author name]]¹ · [[Second author name]]² …', align='left')
    para(doc, '¹ [[Department, University, City, Country]]', align='left')
    para(doc, '² [[Department, University, City, Country]]', align='left', after=10)
    para(doc, '**Corresponding author:** [[Name]], haneenbiate@gmail.com', align='left')
    para(doc, '**ORCID:** [[0000-0000-0000-0000]]', align='left', after=10)
    heading(doc, 'Statements and Declarations', 1)
    para(doc, '**Funding** [[State funding, or: The authors declare that no funds, grants, or other support were '
              'received during the preparation of this manuscript.]]')
    para(doc, '**Competing interests** The authors have no relevant financial or non-financial interests to disclose.')
    para(doc, '**Author contributions** [[CRediT roles for each author.]]')
    para(doc, '**Data availability** Code and text-free artefacts are available at '
              'https://github.com/haneenbiate/- (branch study-v3, tag study-v3-frozen).')
    para(doc, '**Acknowledgements** We thank the three annotators who judged the edited tweets. [[Add others.]]')
    doc.save(output)


def figure_files(output):
    """600-dpi TIFF and PNG for submission, a 300-dpi PNG for embedding, plus the vector PDF."""
    import fitz
    from PIL import Image
    embedded = {}
    for number, stem in FIGURES.items():
        page = fitz.open(str(stem.with_suffix('.pdf')))[0]
        pixmap = page.get_pixmap(dpi=600, alpha=False)
        image = Image.open(io.BytesIO(pixmap.tobytes('png')))
        image.save(output / f'Fig{number}.tif', compression='tiff_lzw', dpi=(600, 600))
        image.save(output / f'Fig{number}.png', dpi=(600, 600))
        shutil.copyfile(stem.with_suffix('.pdf'), output / f'Fig{number}.pdf')
        small = output / f'.embed_fig{number}.png'
        image.resize((image.width // 2, image.height // 2), Image.LANCZOS).save(small, dpi=(300, 300))
        embedded[number] = small
    return embedded


def word_count(path):
    doc = Document(path)
    return sum(len(p.text.split()) for p in doc.paragraphs)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    assert len(ABSTRACT.split()) <= 250, len(ABSTRACT.split())
    data = Data()
    embedded = figure_files(args.output)
    manuscript(data, embedded, args.output / 'Manuscript_anonymized.docx')
    supplementary(data, args.output / 'Online_Resource_1.docx')
    title_page(args.output / 'Title_page.docx')
    for small in embedded.values():
        small.unlink()
    print('abstract words:', len(ABSTRACT.split()), '| manuscript words:',
          word_count(args.output / 'Manuscript_anonymized.docx'))


if __name__ == '__main__':
    main()
